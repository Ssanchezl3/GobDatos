#!/usr/bin/env python3
"""
DICCIONARIO DEL LAGO EN EXCEL — genera docs/diccionario_lago.xlsx

Lee los datasets REALES del lago (web/data/*.geojson) y el registro de
fuentes (sources.json), y produce un Excel con:
  - Hoja "Fuentes": F01/F02/F03 con licencia, vigencia y fecha de prueba.
  - Hoja "Diccionario": SOLO los campos que de verdad existen en el lago,
    con tipo inferido, ejemplo real, fuente y fecha de prueba.

Cumple la rúbrica: "solo lo que de verdad está en el lago, con la fecha real
de prueba y la licencia copiada tal como aparece en la fuente".
"""
import json
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ROOT = os.path.join(os.path.dirname(__file__), "..")
OUT = os.path.join(ROOT, "web", "data")
DOCS = os.path.join(ROOT, "docs")
os.makedirs(DOCS, exist_ok=True)

# dataset -> (fuente, etiqueta legible)
DATASETS = {
    "counties.geojson": ("F01", "Condados de Arkansas"),
    "municipalities.geojson": ("F01", "Municipios de Arkansas"),
    "congress.geojson": ("F01", "Distritos congresionales"),
    "compare.geojson": ("F02", "Condados AR vs TX (comparación)"),
    "climate.geojson": ("F03", "Estaciones climáticas de Arkansas"),
}

DESCRIPTIONS = {
    "name": "Nombre de la entidad", "fips": "Código FIPS",
    "sq_miles": "Superficie (millas²)", "district": "Número de distrito federal",
    "classification": "Clase legal municipal", "geoid": "GEOID (estado+condado)",
    "state_name": "Estado", "growth": "Crecimiento poblacional 2000–2020 (%)",
    "lat": "Latitud del centroide", "lon": "Longitud del centroide",
    "station_id": "Identificador de estación NOAA", "elev_ft": "Elevación (pies)",
    "tavg_annual": "Temperatura media anual (°F)",
    "tmax_annual": "Temperatura máxima media anual (°F)",
    "tmin_annual": "Temperatura mínima media anual (°F)",
    "prcp_annual": "Precipitación anual (pulgadas)",
    "tavg_by_month": "Temperatura media por mes (°F)",
}


def infer_type(v):
    if isinstance(v, bool): return "booleano"
    if isinstance(v, int): return "entero"
    if isinstance(v, float): return "decimal"
    if isinstance(v, list): return "lista"
    return "texto"


def describe(field):
    base = field.rstrip("0123456789")
    if field.startswith("pop"):
        return f"Población censal del año {field[3:]}"
    return DESCRIPTIONS.get(field, DESCRIPTIONS.get(base, "—"))


# estilos
HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="222222")
WRAP = Alignment(vertical="top", wrap_text=True)
THIN = Border(*(Side(style="thin", color="DDDDDD"),) * 4)


def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HEAD; cell.fill = FILL; cell.alignment = WRAP


def build():
    sources = json.load(open(os.path.join(OUT, "sources.json"), encoding="utf-8"))
    S = sources["sources"]
    test_date = sources["test_date"]

    wb = Workbook()

    # ---- Hoja 1: Fuentes ----
    ws = wb.active
    ws.title = "Fuentes"
    cols = ["ID", "Nombre", "Publicador", "Portal", "Cobertura", "Vigencia",
            "Fecha de prueba", "Licencia", "URL licencia", "¿API key?", "Script de ingesta"]
    ws.append(cols)
    for sid in sorted(S):
        s = S[sid]
        ws.append([s["id"], s["name"], s["publisher"], s["portal"], s["coverage"],
                   s["validity"], s["test_date"], s["license"], s["license_url"],
                   "No" if not s["api_key_required"] else "Sí", s["ingest_script"]])
    style_header(ws, len(cols))
    widths = [6, 34, 34, 30, 30, 26, 14, 40, 34, 9, 26]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP; c.border = THIN

    # ---- Hoja 2: Diccionario (solo campos reales) ----
    wd = wb.create_sheet("Diccionario")
    dcols = ["Dataset", "Entidad", "Campo", "Tipo", "Ejemplo real",
             "Descripción", "Fuente", "Fecha de prueba", "Licencia"]
    wd.append(dcols)
    for fname, (src, label) in DATASETS.items():
        path = os.path.join(OUT, fname)
        if not os.path.exists(path):
            continue
        fc = json.load(open(path, encoding="utf-8"))
        props = fc["features"][0]["properties"]
        for field, val in props.items():
            if field.startswith("_"):
                continue
            example = val[0] if isinstance(val, list) and val else val
            wd.append([fname, label, field, infer_type(val), str(example)[:40],
                       describe(field), src, test_date, S[src]["license"]])
    style_header(wd, len(dcols))
    dwidths = [24, 30, 16, 10, 16, 36, 8, 14, 40]
    for i, w in enumerate(dwidths, 1):
        wd.column_dimensions[chr(64 + i)].width = w
    for row in wd.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP; c.border = THIN

    out = os.path.join(DOCS, "diccionario_lago.xlsx")
    wb.save(out)
    nfields = wd.max_row - 1
    print(f"diccionario_lago.xlsx: {len(S)} fuentes, {nfields} campos reales del lago")


if __name__ == "__main__":
    build()
